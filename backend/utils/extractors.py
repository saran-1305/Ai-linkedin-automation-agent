import io
import logging

logger = logging.getLogger(__name__)

def extract_text_from_file(filename: str, file_content: bytes) -> str:
    """Extract text based on the file extension."""
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    
    if ext == "txt":
        return file_content.decode("utf-8", errors="ignore")
        
    elif ext == "pdf":
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                text = "\n".join(page.extract_text() for page in pdf.pages if page.extract_text())
            return text
        except Exception as e:
            logger.error(f"Error extracting PDF {filename}: {e}")
            return ""
            
    elif ext == "docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_content))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
        except Exception as e:
            logger.error(f"Error extracting DOCX {filename}: {e}")
            return ""
            
    elif ext == "pptx":
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_content))
            text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            return "\n".join(text)
        except Exception as e:
            logger.error(f"Error extracting PPTX {filename}: {e}")
            return ""
            
    else:
        logger.warning(f"Unsupported file format: {ext}")
        return ""
