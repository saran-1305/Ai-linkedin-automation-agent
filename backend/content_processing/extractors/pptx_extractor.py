import os
from pptx import Presentation

def extract_pptx(file_path: str) -> tuple[str, dict]:
    """
    Extracts text and metadata from a PPTX file.
    """
    if not file_path or not os.path.exists(file_path):
        raise ValueError(f"Invalid PPTX file path: {file_path}")
        
    prs = Presentation(file_path)
    text_blocks = []
    
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text.append(shape.text)
        
        if slide_text:
            text_blocks.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_text))
            
    extracted_text = "\n\n".join(text_blocks)
    
    metadata = {
        "source_type": "pptx",
        "slide_count": len(prs.slides),
        "filename": os.path.basename(file_path)
    }
    
    if prs.core_properties:
        if prs.core_properties.title:
            metadata["title"] = prs.core_properties.title
        if prs.core_properties.author:
            metadata["author"] = prs.core_properties.author
            
    return extracted_text, metadata
