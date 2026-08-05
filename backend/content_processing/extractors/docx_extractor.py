import os
import docx

def extract_docx(file_path: str) -> tuple[str, dict]:
    """
    Extracts text from a DOCX file.
    Extracts paragraphs, headings, lists, and flattens tables.
    """
    if not file_path or not os.path.exists(file_path):
        raise ValueError(f"Invalid DOCX file path: {file_path}")
        
    doc = docx.Document(file_path)
    text_blocks = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            text_blocks.append(para.text.strip())
            
    # Flatten tables
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_data:
                text_blocks.append(" | ".join(row_data))
                
    extracted_text = "\n\n".join(text_blocks)
    
    metadata = {
        "source_type": "docx",
        "filename": os.path.basename(file_path)
    }
    
    core_props = doc.core_properties
    if core_props.title:
        metadata["title"] = core_props.title
    if core_props.author:
        metadata["author"] = core_props.author
        
    return extracted_text, metadata
